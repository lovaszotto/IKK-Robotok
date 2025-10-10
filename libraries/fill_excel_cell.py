import sys
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
import re

# Set console encoding to UTF-8 to avoid encoding issues on Windows
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except AttributeError:
    # For older Python versions
    pass

def main():
    # Parametererek ellenorzese
    if len(sys.argv) < 5:
        print("HIBA: Nem elegendo parameter!")
        print("Hasznalat:")
        print("  - Koordinata mod: python fill_excel_cell.py <excel_file> <sheet_name> <row> <col> <value>")
        print("  - Cell referencia mod: python fill_excel_cell.py <excel_file> <sheet_name> <cell_ref> <value>")
        sys.exit(1)
    
    excel_file = sys.argv[1]
    sheet_name = sys.argv[2]
    
    # Ellenorizzuk, hogy 4 vagy 5 parametert kaptunk
    if len(sys.argv) == 5:
        # 4 parameter: cell_ref formatum (pl. A1, B3)
        cell_ref = sys.argv[3]
        value = sys.argv[4]
        print(f"Cell referencia mod: {cell_ref} = '{value}'")
    else:
        # 5 parameter: row, col, value formatum
        row = int(sys.argv[3])
        col = sys.argv[4]
        value = sys.argv[5]
        
        # Oszlop konvertalasa
        if col.isdigit():
            col_letter = get_column_letter(int(col))
            cell_ref = f"{col_letter}{row}"
        else:
            cell_ref = f"{col}{row}"
        
        print(f"Koordinata mod: sor {row}, oszlop {col} -> {cell_ref} = '{value}'")
    
    try:
        print(f"Megnyitom az Excel fajlt: {excel_file}")
        wb = openpyxl.load_workbook(excel_file)
        
        # Sheet ellenorzese
        if sheet_name not in wb.sheetnames:
            print(f"HIBA: Sheet '{sheet_name}' nem talalhato!")
            print(f"Elerheto sheet-ek: {wb.sheetnames}")
            sys.exit(1)
        
        ws = wb[sheet_name]
        print(f"Sheet '{sheet_name}' megnyitva")
        
        # Eredeti ertek lekerdezese es uj ertek beallitasa
        old_value = ws[cell_ref].value
        ws[cell_ref] = value
        
        # Fajl mentese retry mechanizmussal
        import time
        max_retries = 3
        for retry in range(max_retries):
            try:
                wb.save(excel_file)
                wb.close()
                break
            except PermissionError as pe:
                if retry < max_retries - 1:
                    print(f"Excel fajl zarva, probalkozas {retry + 1}/{max_retries} - varakozas 2 masodperc...")
                    wb.close()
                    time.sleep(2)
                    # Ujra megnyitas a kovetkezo probalkozashoz
                    if retry < max_retries - 1:
                        wb = openpyxl.load_workbook(excel_file)
                        ws = wb[sheet_name]
                        ws[cell_ref] = value
                else:
                    print(f"KRITIKUS HIBA: Excel fajl nem mentheto {max_retries} probalkozas utan: {str(pe)}")
                    wb.close()
                    # Folytatas a hiba ellenere, ne alljon le a teljes teszt
                    return
        
        print(f"Sikeres frissites:")
        print(f"  Cella: {cell_ref}")
        print(f"  Regi ertek: '{old_value}'")  
        print(f"  Uj ertek: '{value}'")
        print(f"Excel fajl mentve: {excel_file}")
        
    except Exception as e:
        print(f"HIBA: {str(e)}")
        import traceback
        traceback.print_exc()
        # Nem kilepunk hibaval, folytatjuk a teszteket
        return

if __name__ == "__main__":
    main()